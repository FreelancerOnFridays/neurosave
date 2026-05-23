from __future__ import annotations

import asyncio
import logging
import re as _re
import time
from datetime import datetime, timedelta, timezone
from typing import Any, Coroutine
from zoneinfo import ZoneInfo

from aiogram import F, Router
from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup, Message
from aiogram import Bot
from sqlalchemy.ext.asyncio import AsyncSession

from bot.config_store import get_business_connection_id
from db.repositories import user_settings as us_repo
from bot.handlers.ghost import generate_digest_text
from bot.reminder_store import delay_from_iso
from db.engine import get_session
from db.models import Contact, Task
from db.repositories import contacts as contact_repo
from db.repositories import integration_configs as cfg_repo
from db.repositories import ghost as ghost_repo
from db.repositories import messages as msg_repo
from db.repositories import tasks as task_repo
from services.ai import (
    ReminderItem,
    answer_from_context,
    embed_text,
    extract_reminder_from_context,
    extract_tasks_from_message,
    generate_dispatch_message,
    get_style_profile,
    parse_dispatch_command,
    parse_reminder_action,
    transcribe_voice,
)

_RU_LABELS = ("через", "мин", "ч")
_UA_LABELS = ("через", "хв", "год")
_EN_LABELS = ("in", "min", "h")

logger = logging.getLogger(__name__)
router = Router()

_bg_tasks: set[asyncio.Task[None]] = set()

# Pending dispatch when contact not found: owner_id → {alias, text, bcid, session}
_pending_dispatch: dict[int, dict[str, object]] = {}

# Pending email waiting for attachment: owner_id → {to, subject, body}
_pending_email: dict[int, dict[str, object]] = {}

# Pending snooze: owner_id → {task_id, task_description, chat_id, message_id}
_pending_snooze: dict[int, dict[str, object]] = {}

# Pending dispatch task confirmation: owner_id → {extracted, contacts, bcid}
_pending_dispatch_tasks: dict[int, dict[str, Any]] = {}

# Pending send preview (before actual send): owner_id → {contacts, intent, send_text, ...}
_pending_send_preview: dict[int, dict[str, Any]] = {}

# Pending /ask search waiting for contact selection: owner_id → query text
_pending_ask: dict[int, str] = {}

# Ghost activation context: owner_id → {context, lang, tz_name, task, has_auto_off, auto_off_iso}
_ghost_contexts: dict[int, dict[str, object]] = {}

# Pending ghost auto-off time change: owner_id → original_msg_id
_pending_ghost_time: dict[int, int] = {}

# Pending ghost away-message text change: owner_id → original_msg_id
_pending_ghost_text: dict[int, int] = {}

# Pending Gmail reply: owner_id → {from_, thread_id, message_id_header, subject, prompt_msg_id}
_pending_gmail_reply: dict[int, dict[str, Any]] = {}

# Pending file analysis: owner_id → {file_id, filename, mime_type, doc_type}
_pending_file: dict[int, dict[str, str]] = {}


async def _handle_pending_snooze(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
) -> bool:
    pending = _pending_snooze.get(owner_id)
    if pending is None:
        return False
    task_id = int(pending["task_id"])  # type: ignore[arg-type]
    desc = str(pending["task_description"])
    us = await us_repo.get_or_create(session, owner_id)
    try:
        parsed = await extract_reminder_from_context(
            context_text=desc,
            trigger_text=text,
            language=us.language,
            tz_name=us.timezone,
        )
        iso = parsed.reminder_time_iso or parsed.scheduled_at_iso
        if not iso:
            raise ValueError("no time")
        new_time = datetime.fromisoformat(iso)
        if new_time.tzinfo is None:
            new_time = new_time.replace(tzinfo=timezone.utc)
    except Exception:
        return False

    _pending_snooze.pop(owner_id, None)
    await task_repo.set_reminder(session, task_id, new_time)
    try:
        local = new_time.astimezone(ZoneInfo(us.timezone))
        now_local = datetime.now(ZoneInfo(us.timezone))
        is_today = local.date() == now_local.date()
        time_label = local.strftime("%H:%M") if is_today else local.strftime("%d.%m в %H:%M")
        await bot.edit_message_text(
            chat_id=int(pending["chat_id"]),  # type: ignore[arg-type]
            message_id=int(pending["message_id"]),  # type: ignore[arg-type]
            text=f"⏰ Напомню в {time_label}\n<i>{desc}</i>",
            parse_mode="HTML",
            reply_markup=None,
        )
    except Exception:
        pass
    try:
        await message.delete()
    except Exception:
        pass
    return True


async def _show_email_preview(
    message: Message,
    owner_id: int,
    to: str,
    subject: str,
    body: str,
    has_attachment: bool = False,
    queued_file_id: str | None = None,
    queued_filename: str = "",
    queued_mime: str = "",
) -> None:
    """Store composed email in pending and show a draft preview with confirm/edit/cancel buttons."""
    _pending_email[owner_id] = {
        "status": "preview",
        "to": to,
        "subject": subject,
        "body": body,
        "has_attachment": has_attachment,
        "queued_file_id": queued_file_id,
        "queued_filename": queued_filename,
        "queued_mime": queued_mime,
    }
    attach_note = ""
    if queued_file_id and queued_filename:
        attach_note = f"\n\n📎 <i>Вложение: {queued_filename}</i>"
    elif has_attachment:
        attach_note = "\n\n📎 <i>После подтверждения прикрепи файл</i>"

    preview = (
        f"📧 <b>Черновик письма</b>\n\n"
        f"<b>Кому:</b> {to}\n"
        f"<b>Тема:</b> {subject}\n\n"
        f"{body[:600]}{'…' if len(body) > 600 else ''}"
        f"{attach_note}"
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Отправить", callback_data=f"email_send:{owner_id}"),
        InlineKeyboardButton(text="✏️ Изменить", callback_data=f"email_edit:{owner_id}"),
        InlineKeyboardButton(text="❌ Отменить", callback_data=f"email_cancel:{owner_id}"),
    ]])
    await message.answer(preview, parse_mode="HTML", reply_markup=keyboard)


def _replace_event_time(text: str, old_iso: str | None, new_iso: str | None, tz_name: str) -> str:
    if not old_iso or not new_iso:
        return text
    try:
        tz = ZoneInfo(tz_name)
        old_dt = datetime.fromisoformat(old_iso)
        new_dt = datetime.fromisoformat(new_iso)
        if old_dt.tzinfo is None:
            old_dt = old_dt.replace(tzinfo=timezone.utc)
        if new_dt.tzinfo is None:
            new_dt = new_dt.replace(tzinfo=timezone.utc)
        old_loc = old_dt.astimezone(tz)
        new_loc = new_dt.astimezone(tz)
        for fmt in ("{h:02d}:{m}", "{h}:{m}"):
            old_str = fmt.format(h=old_loc.hour, m=old_loc.strftime("%M"))
            new_str = fmt.format(h=new_loc.hour, m=new_loc.strftime("%M"))
            if old_str in text:
                return text.replace(old_str, new_str)
    except Exception:
        pass
    return text


def _format_delay(seconds: float, lang: str) -> str:
    lbl = _UA_LABELS if lang == "ua" else (_RU_LABELS if lang == "ru" else _EN_LABELS)
    total_minutes = int(seconds // 60)
    hours, mins = divmod(total_minutes, 60)
    if hours:
        return f"{lbl[0]} {hours} {lbl[2]} {mins} {lbl[1]}" if mins else f"{lbl[0]} {hours} {lbl[2]}"
    return f"{lbl[0]} {total_minutes} {lbl[1]}"


def _format_time_local(iso: str | None, tz_name: str, lang: str) -> str:
    if not iso:
        return ""
    try:
        dt = datetime.fromisoformat(iso)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        local_dt = dt.astimezone(ZoneInfo(tz_name))
        day = str(local_dt.day)
        hour = str(local_dt.hour)
        minute = local_dt.strftime("%M")
        if lang == "ua":
            return f"{day} {_MONTH_UA_SHORT[local_dt.month]} о {hour}:{minute}"
        if lang == "ru":
            return f"{day} {_MONTH_RU_SHORT[local_dt.month]} в {hour}:{minute}"
        return f"{local_dt.strftime('%b')} {day} at {hour}:{minute}"
    except Exception:
        return iso or ""


def _task_reminder_iso(task: Task) -> str | None:
    return task.reminder_time.isoformat() if task.reminder_time else None


def _task_deadline_iso(task: Task) -> str | None:
    return task.deadline.isoformat() if task.deadline else None


def _build_reminders_ctx(tasks: list[Task], tz_name: str, lang: str) -> str:
    lines = [
        f"{i}. {t.description} — reminder at "
        f"{_format_time_local(_task_reminder_iso(t), tz_name, lang)}"
        + (f", event: {_format_time_local(_task_deadline_iso(t), tz_name, lang)}" if t.deadline else "")
        for i, t in enumerate(tasks, 1)
    ]
    return "\n".join(lines) if lines else "(none)"


def _find_task_by_hint(tasks: list[Task], hint: str | None) -> Task | None:
    if not hint or not tasks:
        return None
    hint_lower = hint.lower()
    for t in tasks:
        if hint_lower in t.description.lower() or t.description.lower() in hint_lower:
            return t
    for t in tasks:
        if any(len(w) > 2 and w in t.description.lower() for w in hint_lower.split()):
            return t
    return tasks[-1] if len(tasks) == 1 else None


def _parse_search_time_range(
    text: str,
    tz_name: str,
) -> tuple[datetime | None, datetime | None]:
    tz = ZoneInfo(tz_name)
    now_local = datetime.now(tz)
    today = now_local.replace(hour=0, minute=0, second=0, microsecond=0)
    tl = text.lower()
    if "сегодня" in tl or "today" in tl:
        return (
            today.astimezone(timezone.utc),
            (today + timedelta(days=1) - timedelta(seconds=1)).astimezone(timezone.utc),
        )
    if "вчера" in tl or "yesterday" in tl:
        yesterday = today - timedelta(days=1)
        return yesterday.astimezone(timezone.utc), (today - timedelta(seconds=1)).astimezone(timezone.utc)
    if "на этой неделе" in tl or "this week" in tl:
        week_start = today - timedelta(days=today.weekday())
        return week_start.astimezone(timezone.utc), now_local.astimezone(timezone.utc)
    if "на прошлой неделе" in tl or "last week" in tl:
        this_week = today - timedelta(days=today.weekday())
        last_week = this_week - timedelta(weeks=1)
        return last_week.astimezone(timezone.utc), (this_week - timedelta(seconds=1)).astimezone(timezone.utc)
    return None, None


def _extract_search_contact(text: str, contacts: list[Contact]) -> Contact | None:
    text_lower = text.lower()

    def _matches(name: str | None) -> bool:
        if not name:
            return False
        name_lower = name.lower()
        if name_lower in text_lower:
            return True
        for part in name_lower.split():
            if len(part) < 3:
                continue
            for token in text_lower.split():
                if len(token) < 3:
                    continue
                min_len = min(len(part), len(token))
                if min_len >= 3 and (
                    part[:min_len] == token[:min_len] or token in part or part in token
                ):
                    return True
        return False

    for c in contacts:
        if _matches(c.saved_name):
            return c
    for c in contacts:
        if _matches(c.name):
            return c
    return None


_PERSON_RE = _re.compile(r'\b(?:с|со|у|от|про)\s+([А-ЯЁа-яё]{3,})', _re.IGNORECASE)


def _normalize_ru_name(word: str) -> str:
    """Strip common Russian oblique-case endings and capitalize."""
    n = word.lower()
    if n.endswith("ем") and len(n) > 4:    # Андреем → Андрей
        n = n[:-2] + "й"
    elif n.endswith("ом") and len(n) > 4:  # Максимом → Максим, Виктором → Виктор
        n = n[:-2]
    elif n.endswith("ем") and len(n) > 3:  # fallback short names
        n = n[:-2] + "й"
    elif n.endswith("ей") and len(n) > 4:  # Машей → Маша
        n = n[:-2] + "а"
    elif n.endswith("ой") and len(n) > 4:  # Светой/Димой → Света/Дима
        n = n[:-2] + "а"
    elif n.endswith("ым") and len(n) > 4:  # Максимым → rare but handle
        n = n[:-2]
    return n.capitalize()


_EXPLICIT_REMINDER_RU = _re.compile(
    r"\bнапомни(те|й)?\b|\bнапоминани[еяю]\b|\bпоставь\s+напоминани[еяю]\b"
    r"|\bдобавь\s+напоминани[еяю]\b|\bне\s+забудь\b",
    _re.IGNORECASE,
)
_EXPLICIT_REMINDER_EN = _re.compile(
    r"\bremind\s+me\b|\bset\s+(a\s+)?reminder\b|\bschedule\s+(a\s+)?reminder\b",
    _re.IGNORECASE,
)
_TASK_KEYWORDS_RU = _re.compile(
    r"\b(нужно|надо|нужна|нужен|должен|должна|должны|необходимо|требуется|следует)\b",
    _re.IGNORECASE,
)


def _is_explicit_reminder(text: str) -> bool:
    return bool(_EXPLICIT_REMINDER_RU.search(text) or _EXPLICIT_REMINDER_EN.search(text))


def _is_task_statement(text: str) -> bool:
    return bool(_TASK_KEYWORDS_RU.search(text))


_MONTH_RU_SHORT = ["", "янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
_MONTH_UA_SHORT = ["", "січ", "лют", "бер", "кві", "тра", "чер", "лип", "сер", "вер", "жов", "лис", "гру"]


_STOPWORDS = frozenset({
    "сегодня", "завтра", "вчера", "утром", "вечером", "ночью", "днём", "дней",
    "через", "после", "перед", "около", "часов", "минут", "today", "tomorrow",
    "monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday",
    "понедельник", "вторник", "среду", "четверг", "пятницу", "субботу", "воскресенье",
    # Generic action verbs — too common to reliably signal a match
    "купить", "сделать", "написать", "позвонить", "отправить", "взять",
    "принести", "забрать", "закончить", "проверить", "посмотреть", "прочитать",
})


def _find_matching_task(tasks: list[Task], new_text: str) -> Task | None:
    new_lower = new_text.lower()
    for t in tasks:
        existing = t.description.lower()
        if existing in new_lower or new_lower in existing:
            return t
        existing_words = {w for w in existing.split() if len(w) > 3 and w not in _STOPWORDS}
        new_words = {w for w in new_lower.split() if len(w) > 3 and w not in _STOPWORDS}
        if not existing_words or not new_words:
            continue
        if existing_words & new_words:
            return t
        for ew in existing_words:
            for nw in new_words:
                stem = min(5, len(ew), len(nw))
                if stem >= 4 and ew[:stem] == nw[:stem]:
                    return t
    return None


def _fire(coro: Coroutine[Any, Any, None]) -> None:
    task = asyncio.create_task(coro)
    _bg_tasks.add(task)
    task.add_done_callback(_bg_tasks.discard)


async def _delayed_send(
    bot: Bot,
    owner_id: int,
    chat_id: int,
    contact_name: str,
    text: str,
    business_connection_id: str | None,
    delay_seconds: float,
) -> None:
    if delay_seconds > 0:
        await asyncio.sleep(delay_seconds)
    try:
        await bot.send_message(chat_id=chat_id, text=text, business_connection_id=business_connection_id)
        await bot.send_message(chat_id=owner_id, text=f"✅ Сообщение отправлено: {contact_name}")
    except Exception:
        logger.exception("Failed to send dispatch message to chat %d", chat_id)
        try:
            await bot.send_message(chat_id=owner_id, text=f"❌ Не удалось отправить сообщение: {contact_name}")
        except Exception:
            pass


async def _ghost_auto_off(bot: Bot, owner_id: int, delay_seconds: float) -> None:
    if delay_seconds > 0:
        await asyncio.sleep(delay_seconds)
    try:
        async with get_session() as session:
            gs = await ghost_repo.get_session(session, owner_id)
            if gs is None or not gs.is_active:
                return  # Already turned off manually
            await ghost_repo.set_active(session, owner_id, active=False)
            await ghost_repo.set_auto_off(session, owner_id, None)
        await bot.send_message(chat_id=owner_id, text="👻 Ghost Mode автоматически выключен.")
    except Exception:
        logger.exception("Failed to auto-deactivate ghost mode for owner %d", owner_id)


def _fire_ghost_auto_off(bot: Bot, owner_id: int, delay: float) -> "asyncio.Task[None]":
    task: asyncio.Task[None] = asyncio.create_task(_ghost_auto_off(bot, owner_id, delay))
    _bg_tasks.add(task)
    task.add_done_callback(_bg_tasks.discard)
    return task


def _build_ghost_status_lines(
    away_text: str,
    lang: str,
    tz_name: str,
    auto_off_iso: str | None,
) -> list[str]:
    lines = ["👻 <b>Ghost Mode включён</b>"]
    if away_text:
        lines.append(f"\n💬 <b>Автоответ:</b>\n<blockquote>{away_text}</blockquote>")
    if auto_off_iso:
        time_label = _format_time_local(auto_off_iso, tz_name, lang)
        lines.append(f"⏰ Авто-выключение: <b>{time_label}</b>")
    else:
        lines.append("ℹ️ Напишите «я свободен» чтобы выключить")
    lines.append("\n<i>Чтобы изменить текст автоответа — напишите его сюда или измените в мини-приложении</i>")
    return lines


def _ghost_activation_keyboard(silent_mode: bool = False) -> InlineKeyboardMarkup:
    silent_label = "🔔 Отправлять автоответ" if silent_mode else "🔕 Не отправлять автоответ"
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="🔄 Сгенерировать", callback_data="ghost_regen"),
            InlineKeyboardButton(text="✏️ Изменить текст", callback_data="ghost_change_text"),
        ],
        [
            InlineKeyboardButton(text="⏰ Изменить время", callback_data="ghost_change_time"),
            InlineKeyboardButton(text=silent_label, callback_data="ghost_toggle_silent"),
        ],
    ])


def _contact_picker_keyboard(alias: str, contacts: list[Contact]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    row: list[InlineKeyboardButton] = []
    for c in contacts:
        label = (c.saved_name or c.name or c.username or str(c.user_id))[:20]
        row.append(InlineKeyboardButton(
            text=label,
            callback_data=f"pick_alias:{alias[:20]}:{c.user_id}",
        ))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    return InlineKeyboardMarkup(inline_keyboard=rows)


@router.message(F.contact)
async def handle_contact_share(message: Message, session: AsyncSession) -> None:
    if message.from_user is None:
        return
    owner_id = message.from_user.id
    c = message.contact
    if not c or not c.user_id:
        await message.answer("Не удалось определить Telegram ID этого контакта.")
        return

    parts = [c.first_name]
    if c.last_name:
        parts.append(c.last_name)
    full_name = " ".join(parts)

    await contact_repo.upsert_contact(session, owner_id=owner_id, user_id=c.user_id,
                                      name=full_name, has_business_chat=True)
    await contact_repo.set_saved_name(session, owner_id=owner_id, user_id=c.user_id,
                                      saved_name=full_name)
    await message.answer(
        f"✅ Контакт сохранён: <b>{full_name}</b>\n"
        f"Теперь можно писать «Напиши {full_name.split()[0]}…»",
        parse_mode="HTML",
    )


def _preview_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Отправить", callback_data="send_preview_send"),
        InlineKeyboardButton(text="✏️ Изменить", callback_data="send_preview_edit"),
        InlineKeyboardButton(text="❌ Отмена", callback_data="send_preview_cancel"),
    ]])


async def _execute_send_from_preview(
    owner_id: int,
    data: dict[str, Any],
    bot: Bot,
    session: AsyncSession,
    conf_message: Message,
) -> None:
    """Execute the confirmed dispatch: send to all contacts, then run task detection."""
    contacts: list[dict[str, Any]] = data["contacts"]
    sent_names: list[str] = []
    dispatched: list[dict[str, Any]] = []

    for c_dict in contacts:
        user_id: int = c_dict["user_id"]
        has_bc: bool = c_dict.get("has_business_chat", False)

        per_msg: str | None = c_dict.get("message")
        if per_msg:
            send_text: str = per_msg
        elif data.get("is_personalized"):
            try:
                send_text = await generate_dispatch_message(
                    intent=data["intent"],
                    recipient_name=c_dict["name"],
                    language=data.get("lang", "ru"),
                    style_profile=data.get("style", ""),
                )
            except Exception:
                send_text = data.get("send_text") or data["intent"]
        else:
            send_text = data["send_text"]

        delay: float = data.get("delay", 0.0)
        bcid: str | None = data.get("bcid")

        if has_bc:
            _fire(_delayed_send(bot, owner_id, user_id, c_dict["name"], send_text, bcid, delay))
        else:
            continue

        sent_names.append(c_dict["name"])
        dispatched.append(c_dict)

    if not dispatched:
        try:
            await conf_message.edit_text("⚠️ Не удалось отправить — нет доступа к чатам.", reply_markup=None)
        except Exception:
            pass
        return

    # Task detection
    extracted_list: list[Any] = []
    try:
        extracted_list = await extract_tasks_from_message(
            data.get("original_text", ""),
            language=data.get("lang", "ru"),
            tz_name=data.get("tz_name", "UTC"),
        )
    except Exception:
        logger.exception("Task detection failed after send")

    if extracted_list:
        _pending_dispatch_tasks[owner_id] = {
            "extracted": extracted_list,
            "contacts": dispatched,
            "bcid": data.get("bcid"),
        }
        first = extracted_list[0]
        desc_preview = (first.description or "")[:80]
        names_str = ", ".join(c["name"] for c in dispatched)
        deadline_line = ""
        if first.deadline_iso:
            try:
                dl = datetime.fromisoformat(first.deadline_iso)
                deadline_line = f"\n⏰ Дедлайн: {dl.strftime('%d.%m.%Y')}"
            except ValueError:
                pass
        try:
            await conf_message.edit_text(
                f"✅ Отправлено: {', '.join(sent_names)}\n\n"
                f"📋 <b>Добавить в делегированные задачи?</b>\n"
                f"Задача: {desc_preview}\n"
                f"Исполнитель(и): {names_str}{deadline_line}",
                parse_mode="HTML",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                    InlineKeyboardButton(text="✅ Да", callback_data="dispatch_task_yes"),
                    InlineKeyboardButton(text="❌ Нет", callback_data="dispatch_task_no"),
                ]]),
            )
        except Exception:
            pass
    else:
        try:
            await conf_message.edit_text(
                f"✅ Отправлено: {', '.join(sent_names)}",
                reply_markup=None,
            )
        except Exception:
            pass


async def _handle_pending_ghost_time(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
) -> bool:
    if owner_id not in _pending_ghost_time:
        return False
    original_msg_id = _pending_ghost_time.pop(owner_id)
    ctx = _ghost_contexts.get(owner_id, {})
    lang = str(ctx.get("lang", "ru"))
    tz_name = str(ctx.get("tz_name", "Europe/Moscow"))

    try:
        parsed2 = await extract_reminder_from_context(
            context_text="ghost mode auto-off",
            trigger_text=text,
            language=lang,
            tz_name=tz_name,
        )
        iso: str | None = parsed2.reminder_time_iso or parsed2.scheduled_at_iso
    except Exception:
        iso = None

    if not iso:
        # Re-ask
        _pending_ghost_time[owner_id] = original_msg_id
        try:
            await bot.edit_message_text(
                chat_id=owner_id,
                message_id=original_msg_id,
                text="❌ Не удалось распознать время. Попробуйте: «через 30 минут», «в 19:30»",
                parse_mode="HTML",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                    InlineKeyboardButton(text="❌ Отмена", callback_data="ghost_change_time_cancel"),
                ]]),
            )
        except Exception:
            pass
        try:
            await message.delete()
        except Exception:
            pass
        return True

    # Cancel old timer
    old_task = ctx.get("task")
    if isinstance(old_task, asyncio.Task) and not old_task.done():
        old_task.cancel()

    d = delay_from_iso(iso)
    new_task = _fire_ghost_auto_off(bot, owner_id, d)
    _ghost_contexts[owner_id] = {**ctx, "task": new_task, "has_auto_off": True, "auto_off_iso": iso}

    try:
        auto_off_dt = datetime.fromisoformat(iso)
        if auto_off_dt.tzinfo is None:
            auto_off_dt = auto_off_dt.replace(tzinfo=timezone.utc)
        await ghost_repo.set_auto_off(session, owner_id, auto_off_dt)
    except Exception:
        pass

    gs = await ghost_repo.get_session(session, owner_id)
    away_text = (gs.away_message or "") if gs else ""
    silent_mode = bool(ctx.get("silent_mode", False))
    lines = _build_ghost_status_lines(away_text, lang, tz_name, iso)
    try:
        await bot.edit_message_text(
            chat_id=owner_id,
            message_id=original_msg_id,
            text="\n".join(lines),
            parse_mode="HTML",
            reply_markup=_ghost_activation_keyboard(silent_mode),
        )
    except Exception:
        pass
    try:
        await message.delete()
    except Exception:
        pass
    return True


async def _handle_pending_ghost_text(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
) -> bool:
    if owner_id not in _pending_ghost_text:
        return False
    original_msg_id = _pending_ghost_text.pop(owner_id)
    new_away = text.strip()
    if not new_away:
        return False

    await ghost_repo.update_away_message(session, owner_id, new_away)

    ctx = _ghost_contexts.get(owner_id, {})
    lang = str(ctx.get("lang", "ru"))
    tz_name = str(ctx.get("tz_name", "Europe/Moscow"))
    auto_off_iso = ctx.get("auto_off_iso")
    silent_mode = bool(ctx.get("silent_mode", False))

    lines = _build_ghost_status_lines(
        new_away, lang, tz_name,
        str(auto_off_iso) if auto_off_iso else None,
    )
    try:
        await bot.edit_message_text(
            chat_id=owner_id,
            message_id=original_msg_id,
            text="\n".join(lines),
            parse_mode="HTML",
            reply_markup=_ghost_activation_keyboard(silent_mode),
        )
    except Exception:
        pass
    try:
        await message.delete()
    except Exception:
        pass
    return True


async def _handle_pending_gmail_reply(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
) -> bool:
    pending = _pending_gmail_reply.get(owner_id)
    if pending is None:
        return False

    _pending_gmail_reply.pop(owner_id, None)

    prompt_msg_id = pending.get("prompt_msg_id")
    if prompt_msg_id:
        try:
            await bot.delete_message(chat_id=owner_id, message_id=int(prompt_msg_id))  # type: ignore[arg-type]
        except Exception:
            pass

    from services import gmail as gmail_svc

    gmail_service = await gmail_svc.get_gmail_service(owner_id, session)
    if gmail_service is None:
        await message.answer("❌ Gmail не подключён.")
        return True

    from_ = str(pending.get("from_", ""))
    import re as _reply_re
    email_match = _reply_re.search(r"<([^>]+)>", from_)
    to_email = email_match.group(1) if email_match else from_.strip()

    subject = str(pending.get("subject", ""))
    reply_subject = subject if subject.lower().startswith("re:") else f"Re: {subject}"
    thread_id = str(pending.get("thread_id", "")) or None
    in_reply_to = str(pending.get("message_id_header", "")) or None

    try:
        await gmail_svc.send_reply(
            gmail_service,
            to=[to_email],
            subject=reply_subject,
            body=text.strip(),
            thread_id=thread_id,
            in_reply_to=in_reply_to,
        )
        await message.answer(
            f"✅ Ответ отправлен на <b>{to_email}</b>",
            parse_mode="HTML",
        )
    except Exception:
        logger.exception("Gmail reply send failed for owner %d", owner_id)
        await message.answer("❌ Не удалось отправить ответ. Попробуйте ещё раз.")
    return True


async def _process_owner_text(
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
    owner_id: int,
) -> None:
    us = await us_repo.get_or_create(session, owner_id)
    lang = us.language
    tz_name = us.timezone
    original_text = text

    # ── Edit-mode: user typed new message text after clicking "Изменить" ──────
    psp = _pending_send_preview.get(owner_id)
    if psp and psp.get("edit_mode"):
        psp["send_text"] = text
        psp["edit_mode"] = False
        psp["is_personalized"] = False
        conf_msg_id = psp.get("conf_msg_id")
        if conf_msg_id:
            try:
                await bot.delete_message(chat_id=owner_id, message_id=int(conf_msg_id))
            except Exception:
                pass
        names = ", ".join(c["name"] for c in psp["contacts"])
        await message.answer(
            f"📤 <b>{names}</b> будет отправлено:\n\n«{text}»\n\nВсё верно?",
            parse_mode="HTML",
            reply_markup=_preview_keyboard(),
        )
        return

    parsed = await parse_dispatch_command(text, language=lang, tz_name=tz_name)

    # ── Settings change ───────────────────────────────────────────────────────
    if parsed.is_settings and parsed.timezone_iana:
        try:
            from zoneinfo import ZoneInfoNotFoundError
            ZoneInfo(parsed.timezone_iana)
            await us_repo.update_settings(session, owner_id, timezone=parsed.timezone_iana)
            await message.answer(f"✅ Часовой пояс изменён: <b>{parsed.timezone_iana}</b>", parse_mode="HTML")
        except (KeyError, ZoneInfoNotFoundError):
            await message.answer(
                f"❌ Не удалось распознать часовой пояс «{parsed.timezone_iana}». "
                "Попробуйте формат: <code>Europe/Warsaw</code>", parse_mode="HTML",
            )
        return

    # ── Ghost Mode ────────────────────────────────────────────────────────────
    if parsed.is_ghost:
        if parsed.ghost_active:
            gs = await ghost_repo.get_session(session, owner_id)
            if gs is not None and gs.is_active and parsed.ghost_away_message:
                await ghost_repo.update_away_message(session, owner_id, parsed.ghost_away_message)
                await message.answer(f"✅ Сообщение обновлено: «{parsed.ghost_away_message}»")
            else:
                gs_result = await ghost_repo.set_active(session, owner_id, active=True,
                                                       away_message=parsed.ghost_away_message)
                silent_mode = gs_result.silent_mode or False
                task_ref: asyncio.Task[None] | None = None
                if parsed.ghost_until_iso:
                    d = delay_from_iso(parsed.ghost_until_iso)
                    task_ref = _fire_ghost_auto_off(bot, owner_id, d)
                    try:
                        auto_off_dt = datetime.fromisoformat(parsed.ghost_until_iso)
                        if auto_off_dt.tzinfo is None:
                            auto_off_dt = auto_off_dt.replace(tzinfo=timezone.utc)
                        await ghost_repo.set_auto_off(session, owner_id, auto_off_dt)
                    except Exception:
                        pass
                else:
                    await ghost_repo.set_auto_off(session, owner_id, None)
                _ghost_contexts[owner_id] = {
                    "context": original_text,
                    "lang": lang,
                    "tz_name": tz_name,
                    "task": task_ref,
                    "has_auto_off": bool(parsed.ghost_until_iso),
                    "auto_off_iso": parsed.ghost_until_iso,
                    "silent_mode": silent_mode,
                }
                lines = _build_ghost_status_lines(
                    parsed.ghost_away_message or "",
                    lang, tz_name,
                    parsed.ghost_until_iso,
                )
                await message.answer(
                    "\n".join(lines),
                    parse_mode="HTML",
                    reply_markup=_ghost_activation_keyboard(silent_mode),
                )
        else:
            digest = await generate_digest_text(session, owner_id)
            await ghost_repo.set_active(session, owner_id, active=False)
            _ghost_contexts.pop(owner_id, None)
            await message.answer("👻 Ghost Mode выключен. Добро пожаловать обратно!")
            if digest:
                await message.answer(digest, parse_mode="HTML")
        return

    # ── Personal reminder ─────────────────────────────────────────────────────
    if parsed.is_reminder and (parsed.reminder_items or parsed.reminder_text):
        # ── Task statement without explicit reminder keyword → create/update task ─
        # e.g. "нужно сегодня закончить проект" or "огурцы надо купить в 20:00"
        if (
            _is_task_statement(original_text)
            and not _is_explicit_reminder(original_text)
            and len(parsed.reminder_items) < 2
        ):
            task_desc = parsed.reminder_text or original_text.strip()

            # Use AI-parsed times when available (e.g. "надо купить в 20:00")
            stmt_reminder: datetime | None = None
            if parsed.reminder_time_iso:
                try:
                    stmt_reminder = datetime.fromisoformat(parsed.reminder_time_iso)
                    if stmt_reminder.tzinfo is None:
                        stmt_reminder = stmt_reminder.replace(tzinfo=timezone.utc)
                except ValueError:
                    pass

            stmt_deadline: datetime | None = None
            if parsed.event_time_iso:
                try:
                    stmt_deadline = datetime.fromisoformat(parsed.event_time_iso)
                    if stmt_deadline.tzinfo is None:
                        stmt_deadline = stmt_deadline.replace(tzinfo=timezone.utc)
                except ValueError:
                    pass

            # Only fall back to today midnight when no time was given at all
            if stmt_deadline is None and stmt_reminder is None:
                tz = ZoneInfo(tz_name)
                today_local_ts = datetime.now(tz).replace(hour=0, minute=0, second=0, microsecond=0)
                stmt_deadline = today_local_ts.astimezone(timezone.utc)

            # Update existing matching task instead of creating a duplicate
            active_tasks_stmt = await task_repo.get_open_tasks(session, owner_id, is_personal=True)
            existing_stmt = _find_matching_task(active_tasks_stmt, task_desc)

            time_iso_stmt = parsed.reminder_time_iso or parsed.event_time_iso
            time_str_stmt = _format_time_local(time_iso_stmt, tz_name, lang) if time_iso_stmt else ""

            if existing_stmt is not None:
                if stmt_deadline is not None:
                    existing_stmt.deadline = stmt_deadline
                await task_repo.set_reminder(session, existing_stmt.id, stmt_reminder)
                if time_str_stmt:
                    await message.answer(
                        f"✅ Задача обновлена: {existing_stmt.description}\n⏰ Напомню в {time_str_stmt}"
                    )
                else:
                    await message.answer(f"✅ Задача обновлена: {existing_stmt.description}")
            else:
                new_task = await task_repo.create_personal_task(
                    session, owner_id=owner_id, description=task_desc,
                    deadline=stmt_deadline, reminder_time=stmt_reminder,
                )
                if time_str_stmt:
                    await message.answer(
                        f"✅ Задача добавлена: {new_task.description}\n⏰ Напомню в {time_str_stmt}"
                    )
                else:
                    tz = ZoneInfo(tz_name)
                    today_local = datetime.now(tz)
                    if lang == "ua":
                        date_str = f"{today_local.day} {_MONTH_UA_SHORT[today_local.month]}"
                    elif lang == "ru":
                        date_str = f"{today_local.day} {_MONTH_RU_SHORT[today_local.month]}"
                    else:
                        date_str = f"{today_local.strftime('%b')} {today_local.day}"
                    await message.answer(
                        f"✅ Задача добавлена: {new_task.description}\n📅 {date_str}"
                    )
            return

        # ── Multiple reminders: batch create ──────────────────────────────────
        if len(parsed.reminder_items) >= 2:
            confirms: list[str] = []
            for item in parsed.reminder_items:
                r_time: datetime | None = None
                if item.reminder_time_iso:
                    try:
                        r_time = datetime.fromisoformat(item.reminder_time_iso)
                        if r_time.tzinfo is None:
                            r_time = r_time.replace(tzinfo=timezone.utc)
                    except ValueError:
                        pass
                e_time: datetime | None = None
                if item.event_time_iso:
                    try:
                        e_time = datetime.fromisoformat(item.event_time_iso)
                        if e_time.tzinfo is None:
                            e_time = e_time.replace(tzinfo=timezone.utc)
                    except ValueError:
                        pass
                created_t = await task_repo.create_personal_task(
                    session, owner_id=owner_id,
                    description=item.reminder_text,
                    deadline=e_time,
                    reminder_time=r_time,
                )
                time_str = _format_time_local(item.reminder_time_iso, tz_name, lang) if item.reminder_time_iso else ""
                confirms.append(f"• {created_t.description}" + (f" — {time_str}" if time_str else ""))
            await message.answer(
                f"✅ Создано {len(confirms)} напоминания:\n" + "\n".join(confirms)
                + "\nЧтобы изменить или удалить — напишите мне."
            )
            return

        # ── Single reminder: existing logic ───────────────────────────────────
        iso = parsed.reminder_time_iso or parsed.scheduled_at_iso

        reminder_time: datetime | None = None
        if iso:
            try:
                reminder_time = datetime.fromisoformat(iso)
                if reminder_time.tzinfo is None:
                    reminder_time = reminder_time.replace(tzinfo=timezone.utc)
            except ValueError:
                pass

        deadline: datetime | None = None
        if parsed.event_time_iso:
            try:
                deadline = datetime.fromisoformat(parsed.event_time_iso)
                if deadline.tzinfo is None:
                    deadline = deadline.replace(tzinfo=timezone.utc)
            except ValueError:
                pass

        active_tasks = await task_repo.get_open_tasks(session, owner_id, is_personal=True)
        existing_task = _find_matching_task(active_tasks, parsed.reminder_text)

        if existing_task is None and active_tasks and not _is_explicit_reminder(original_text):
            ctx = _build_reminders_ctx(active_tasks, tz_name, lang)
            try:
                action = await parse_reminder_action(text, ctx, language=lang, tz_name=tz_name)
            except Exception:
                action = None
            if action and action.action == "adjust_time" and action.new_reminder_time_iso:
                target = _find_task_by_hint(active_tasks, action.reminder_hint)
                if target:
                    new_reminder_time: datetime | None = None
                    try:
                        new_reminder_time = datetime.fromisoformat(action.new_reminder_time_iso)
                        if new_reminder_time.tzinfo is None:
                            new_reminder_time = new_reminder_time.replace(tzinfo=timezone.utc)
                    except ValueError:
                        pass
                    new_desc = _replace_event_time(target.description, _task_deadline_iso(target),
                                                   parsed.event_time_iso, tz_name)
                    target.description = new_desc
                    await task_repo.set_reminder(session, target.id, new_reminder_time)
                    time_str = _format_time_local(action.new_reminder_time_iso, tz_name, lang)
                    confirm = (
                        f"✅ Напоминание обновлено: {new_desc}\n"
                        f"⏰ Напомню {action.lead_description} ({time_str})"
                        if action.lead_description
                        else f"✅ Напоминание обновлено: {new_desc}\n⏰ Напомню в {time_str}"
                    )
                    await message.answer(confirm)
                    return

        is_update = existing_task is not None
        if is_update and existing_task is not None:
            new_desc = _replace_event_time(existing_task.description, _task_deadline_iso(existing_task),
                                           parsed.event_time_iso, tz_name)
            existing_task.description = new_desc
            if deadline is not None:
                existing_task.deadline = deadline
            await task_repo.set_reminder(session, existing_task.id, reminder_time)
            reminder_text = new_desc
        else:
            new_task = await task_repo.create_personal_task(
                session, owner_id=owner_id, description=parsed.reminder_text,
                deadline=deadline, reminder_time=reminder_time,
            )
            reminder_text = new_task.description

        time_str = _format_time_local(iso, tz_name, lang) if iso else ""
        action_word = "обновлено" if is_update else "добавлено"
        footer = "\nЧтобы изменить время или удалить — просто напишите мне."
        if time_str:
            reminder_line = (
                f"⏰ Напомню {parsed.lead_description} ({time_str})"
                if parsed.lead_description
                else f"⏰ Напомню в {time_str}"
            )
            confirm = f"✅ Напоминание {action_word}: {reminder_text}\n{reminder_line}{footer}"
        else:
            confirm = f"✅ Напоминание {action_word}: {reminder_text}.{footer}"
        await message.answer(confirm)
        return

    # ── Email ─────────────────────────────────────────────────────────────────
    if parsed.is_email and parsed.recipients:
        from services import gmail as gmail_svc
        from services.ai import generate_dispatch_message

        gmail_service = await gmail_svc.get_gmail_service(owner_id, session)
        if gmail_service is None:
            await message.answer(
                "❌ Gmail не подключён. Перейди в мини-приложение → Настройки → Интеграции → Gmail."
            )
            return

        recipient_name = parsed.recipients[0]

        # If recipient looks like a direct email address — use it immediately
        email_direct = _re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", recipient_name.strip())
        if email_direct:
            email_addr: str | None = recipient_name.strip()
            contact = None
        else:
            contacts = await contact_repo.find_contacts_by_name(session, owner_id, recipient_name)
            contact = contacts[0] if contacts else None

            # Email from contact record, fallback to saved cfg mapping
            email_addr = (contact.email or None) if contact else None
            if not email_addr:
                email_addr = await cfg_repo.get_config(session, owner_id, f"email_for:{recipient_name.lower()}")

        if not email_addr:
            _pending_email[owner_id] = {
                "recipient_name": recipient_name,
                "subject": parsed.email_subject or "",
                "body_intent": parsed.email_body_intent,
                "literal_body": parsed.email_literal_body,
                "has_attachment": bool(parsed.email_has_attachment),
                "contact_id": contact.user_id if contact else None,
            }
            name_display = (contact.name or recipient_name) if contact else recipient_name
            await message.answer(
                f"📧 Для {name_display} нет email-адреса в базе.\n"
                "Введи адрес электронной почты:"
            )
            return

        body = parsed.email_literal_body or ""
        if not body:
            style = await get_style_profile(owner_id, [])
            body = await generate_dispatch_message(
                intent=parsed.email_body_intent or "",
                recipient_name=(contact.name if contact else None) or recipient_name,
                language=lang,
                style_profile=style,
            )
        subject = parsed.email_subject or f"Сообщение от {await gmail_svc.get_gmail_address(owner_id, session) or 'NeuroSave'}"

        await _show_email_preview(
            message, owner_id, email_addr, subject, body,
            has_attachment=bool(parsed.email_has_attachment),
        )
        return

    # ── Dispatch to contacts ──────────────────────────────────────────────────
    if parsed.has_dispatch and parsed.recipients:
        if not (parsed.literal_message or parsed.message_intent):
            return

        bcid = get_business_connection_id()
        dispatched_ids: set[int] = set()
        resolved_contacts: list[Contact] = []
        not_found: list[str] = []

        recent_msgs = await msg_repo.get_recent_owner_messages(session, owner_id)
        style = await get_style_profile(owner_id, [m.text for m in recent_msgs])
        dispatch_delay = delay_from_iso(parsed.scheduled_at_iso)

        # ── Phase 1: resolve recipients (no sending) ─────────────────────────
        for recipient_name in parsed.recipients:
            matches = await contact_repo.find_contacts_by_name(
                session, owner_id=owner_id, name=recipient_name,
            )

            if not matches:
                group_contacts = await contact_repo.find_contacts_by_label(
                    session, owner_id=owner_id, label=recipient_name,
                )

                if group_contacts:
                    for gc in group_contacts:
                        if gc.user_id not in dispatched_ids:
                            dispatched_ids.add(gc.user_id)
                            resolved_contacts.append(gc)
                    continue

                not_found.append(recipient_name)
                continue

            contact = matches[0]
            if contact.user_id not in dispatched_ids:
                dispatched_ids.add(contact.user_id)
                resolved_contacts.append(contact)

        # ── Handle not-found with contact picker (unchanged) ─────────────────
        for alias in not_found:
            recent_c = await contact_repo.get_recent_contacts(session, owner_id, limit=12)
            if not recent_c:
                await message.answer(f"❓ Не нашёл «{alias}» — перешлите боту любое сообщение от этого контакта, чтобы я его запомнил.")
                continue
            intent_for_pending = parsed.literal_message or parsed.message_intent or ""
            try:
                pending_text: str = await generate_dispatch_message(
                    intent=intent_for_pending,
                    recipient_name=alias, language=lang, style_profile=style,
                )
            except Exception:
                pending_text = intent_for_pending
            _pending_dispatch[owner_id] = {
                "alias": alias,
                "text": pending_text,
                "business_connection_id": bcid,
            }
            await message.answer(
                f"❓ Не нашёл «{alias}» среди контактов.\n"
                "Выберите кого вы имеете в виду — запомню псевдоним и отправлю:",
                reply_markup=_contact_picker_keyboard(alias, recent_c),
            )

        if not resolved_contacts:
            return

        # ── Phase 2: build per-recipient message map (if different messages per recipient) ──
        # Map contact display_name → literal message from parsed.recipient_messages
        per_recipient_map: dict[str, str] = {}
        if parsed.recipient_messages:
            for rm in parsed.recipient_messages:
                per_recipient_map[rm.recipient.lower()] = rm.message

        def _match_per_msg(contact: "Contact") -> str | None:
            display = (contact.saved_name or contact.name or "").lower()
            for key, msg in per_recipient_map.items():
                if key in display or display in key:
                    return msg
            return None

        # ── Phase 3: generate preview & show confirmation ────────────────────
        intent = parsed.literal_message or parsed.message_intent or ""
        is_personalized = bool(parsed.message_intent) and len(resolved_contacts) > 1

        # Build contact dicts with optional per-recipient message
        contact_dicts: list[dict[str, Any]] = []
        for c in resolved_contacts:
            c_name = c.saved_name or c.name or ""
            c_dict_entry: dict[str, Any] = {
                "user_id": c.user_id,
                "name": c_name,
                "username": c.username or "",
                "has_business_chat": c.has_business_chat,
            }
            pm = _match_per_msg(c)
            if pm:
                c_dict_entry["message"] = pm
            contact_dicts.append(c_dict_entry)

        # Generate preview text from first contact (or first per-recipient message)
        has_per_recipient = bool(per_recipient_map) and any("message" in cd for cd in contact_dicts)
        if has_per_recipient:
            preview_text = contact_dicts[0].get("message") or intent
        else:
            first_name = resolved_contacts[0].saved_name or resolved_contacts[0].name or ""
            try:
                preview_text = await generate_dispatch_message(
                    intent=intent,
                    recipient_name=first_name,
                    language=lang,
                    style_profile=style,
                )
            except Exception:
                preview_text = intent

        _pending_send_preview[owner_id] = {
            "contacts": contact_dicts,
            "intent": intent,
            "send_text": preview_text,
            "is_personalized": is_personalized and not has_per_recipient,
            "bcid": bcid,
            "delay": dispatch_delay,
            "lang": lang,
            "style": style,
            "original_text": original_text,
            "tz_name": tz_name,
            "edit_mode": False,
        }

        delay_note = f" через {_format_delay(dispatch_delay, lang)}" if dispatch_delay > 0 else ""

        if has_per_recipient:
            lines = "\n".join(
                f"📤 <b>{cd['name']}</b>: «{cd.get('message', preview_text)}»"
                for cd in contact_dicts
            )
            await message.answer(
                f"{lines}{delay_note}\n\nОтправить?",
                parse_mode="HTML",
                reply_markup=_preview_keyboard(),
            )
        else:
            names_display = ", ".join(c.saved_name or c.name or str(c.user_id) for c in resolved_contacts)
            personalized_note = "\n<i>(сообщение персонализировано для каждого)</i>" if is_personalized else ""
            await message.answer(
                f"📤 <b>{names_display}</b>{delay_note} будет отправлено:{personalized_note}\n\n"
                f"«{preview_text}»\n\n"
                f"Хотите изменить?",
                parse_mode="HTML",
                reply_markup=_preview_keyboard(),
            )
        return

    # ── Reminder action (adjust / delete) ────────────────────────────────────
    active_tasks = await task_repo.get_open_tasks(session, owner_id, is_personal=True)
    if active_tasks:
        ctx = _build_reminders_ctx(active_tasks, tz_name, lang)
        action = await parse_reminder_action(text, ctx, language=lang, tz_name=tz_name)

        if action.action == "delete":
            target = _find_task_by_hint(active_tasks, action.reminder_hint)
            if target:
                await task_repo.delete_task(session, target.id)
                await message.answer(f"🗑 Напоминание «{target.description}» удалено.")
            else:
                await message.answer("Не нашёл такое напоминание. Используйте /reminders чтобы увидеть список.")
            return

        if action.action == "adjust_time" and action.new_reminder_time_iso:
            target = _find_task_by_hint(active_tasks, action.reminder_hint)
            if target:
                new_rt: datetime | None = None
                try:
                    new_rt = datetime.fromisoformat(action.new_reminder_time_iso)
                    if new_rt.tzinfo is None:
                        new_rt = new_rt.replace(tzinfo=timezone.utc)
                except ValueError:
                    pass
                await task_repo.set_reminder(session, target.id, new_rt)
                time_str = _format_time_local(action.new_reminder_time_iso, tz_name, lang)
                confirm = (
                    f"✅ Напоминание обновлено: {target.description}\n"
                    f"⏰ Напомню {action.lead_description} ({time_str})"
                    if action.lead_description
                    else f"✅ Напоминание обновлено: {target.description}\n⏰ Напомню в {time_str}"
                )
                await message.answer(confirm)
            else:
                await message.answer("Не нашёл такое напоминание. Используйте /reminders чтобы увидеть список.")
            return

    # ── Semantic search — AI-detected or keyword-triggered ────────────────────
    _SEARCH_KW = ("ищи ", "ищи,", "ищи.", "найди ", "найди,", "о чем ", "о чём ", "поищи ")
    text_lower = text.lower()
    keyword_search = any(text_lower.startswith(kw) or f" {kw}" in text_lower for kw in _SEARCH_KW)
    if not parsed.is_search and not keyword_search:
        return

    thinking = await message.answer("🔍 Ищу в истории переписок…")
    try:
        since, until = _parse_search_time_range(text, tz_name)
        time_range_explicit = since is not None

        all_contacts = await contact_repo.get_recent_contacts(session, owner_id, limit=500)
        matched_contact = _extract_search_contact(text, all_contacts)

        if matched_contact is None:
            m = _PERSON_RE.search(text)
            if m:
                person_alias = _normalize_ru_name(m.group(1))
                recent_contacts = await contact_repo.get_recent_contacts(session, owner_id, limit=12)
                if recent_contacts:
                    _pending_ask[owner_id] = text
                    await thinking.delete()
                    await message.answer(
                        f"❓ Не нашёл контакт «{person_alias}» — выберите кого вы имеете в виду:\n"
                        "Запомню псевдоним и выполню поиск.",
                        reply_markup=_contact_picker_keyboard(person_alias, recent_contacts),
                    )
                    return

        results: list = []

        if matched_contact is not None:
            effective_since = since or (datetime.now(timezone.utc) - timedelta(days=14))
            effective_until = until or datetime.now(timezone.utc)
            results = await msg_repo.get_messages_in_chat(
                session, owner_id, matched_contact.user_id,
                since=effective_since, until=effective_until, limit=80,
            )
            if not results and time_range_explicit:
                results = await msg_repo.get_messages_in_chat(
                    session, owner_id, matched_contact.user_id, limit=60,
                )

        if not results:
            query_vec = await embed_text(text)
            results = await msg_repo.search_similar(session, owner_id, query_vec, limit=20)

        if not results:
            await thinking.edit_text("Не нашёл ничего похожего в истории переписок.")
            return

        name_map = await contact_repo.get_name_map(session, owner_id)
        ans = await answer_from_context(text, results, language=lang, name_map=name_map, tz_name=tz_name)
        await thinking.edit_text(ans, parse_mode="HTML")
    except Exception:
        logger.exception("Semantic search fallback failed")
        await thinking.edit_text("❌ Не удалось выполнить поиск.")


_SHEETS_URL_RE = _re.compile(r"docs\.google\.com/spreadsheets/d/([a-zA-Z0-9_-]+)")
_DOCS_URL_RE = _re.compile(r"docs\.google\.com/document/d/([a-zA-Z0-9_-]+)")

_SPREADSHEET_MIMES = {
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/csv",
    "application/csv",
    "text/comma-separated-values",
}
_SPREADSHEET_EXTS = {".csv", ".xlsx", ".xls"}


def _is_spreadsheet_file(filename: str, mime: str) -> bool:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    return mime in _SPREADSHEET_MIMES or ext in _SPREADSHEET_EXTS


_DOCUMENT_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "application/json",
}
_DOCUMENT_EXTS = {".pdf", ".docx", ".doc", ".txt", ".json", ".md"}


def _is_document_file(filename: str, mime: str) -> bool:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    return mime in _DOCUMENT_MIMES or ext in _DOCUMENT_EXTS


def _is_analyzable_file(filename: str, mime: str) -> bool:
    return _is_spreadsheet_file(filename, mime) or _is_document_file(filename, mime)


def _get_doc_type(filename: str, mime: str) -> str:
    return "spreadsheet" if _is_spreadsheet_file(filename, mime) else "document"


def _parse_spreadsheet_bytes(data: bytes, filename: str, mime: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if mime == "text/csv" or ext == "csv":
        import csv, io
        try:
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
            rows = list(reader)
            return "\n".join("\t".join(r) for r in rows[:300])
        except Exception:
            return data.decode("utf-8-sig", errors="replace")[:14000]
    if ext in ("xlsx", "xls") or "spreadsheet" in mime:
        try:
            import openpyxl, io as _io
            wb = openpyxl.load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            if ws is None:
                return ""
            rows_out: list[str] = []
            for row in ws.iter_rows(max_row=300, values_only=True):
                rows_out.append("\t".join("" if c is None else str(c) for c in row))
            return "\n".join(rows_out)
        except Exception as e:
            logger.warning("openpyxl failed for %s: %s", filename, e)
            return ""

    # Plain text / Markdown
    if mime in ("text/plain",) or ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace")[:14000]

    # JSON
    if mime == "application/json" or ext == "json":
        import json as _json
        try:
            obj = _json.loads(data.decode("utf-8", errors="replace"))
            return _json.dumps(obj, ensure_ascii=False, indent=2)[:14000]
        except Exception:
            return data.decode("utf-8", errors="replace")[:14000]

    # PDF
    if mime == "application/pdf" or ext == "pdf":
        try:
            import pypdf as _pypdf
            import io as _pdf_io
            reader = _pypdf.PdfReader(_pdf_io.BytesIO(data))
            pages_text = [page.extract_text() or "" for page in reader.pages[:30]]
            return "\n\n".join(t for t in pages_text if t)[:14000]
        except ImportError:
            logger.warning("pypdf not installed, cannot parse PDF")
            return ""
        except Exception as exc:
            logger.warning("PDF parse failed for %s: %s", filename, exc)
            return ""

    # DOCX
    if ext == "docx" or "wordprocessingml" in mime:
        try:
            import docx as _docx
            import io as _docx_io
            doc = _docx.Document(_docx_io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text)[:14000]
        except ImportError:
            logger.warning("python-docx not installed, cannot parse DOCX")
            return ""
        except Exception as exc:
            logger.warning("DOCX parse failed for %s: %s", filename, exc)
            return ""

    return ""


async def _analyze_uploaded_file(
    message: Message,
    bot: "Bot",
    owner_id: int,
    file_id: str,
    filename: str,
    mime_type: str,
    doc_type: str = "spreadsheet",
) -> None:
    from services.ai import analyze_document

    question = message.caption or "Проанализируй содержимое, дай краткую сводку"
    await message.answer("🔍 Читаю файл…")

    try:
        buf = await bot.download(file_id)
        if buf is None:
            await message.answer("❌ Не удалось скачать файл.")
            return
        data = buf.read()
    except Exception:
        logger.exception("Failed to download file for owner %d", owner_id)
        await message.answer("❌ Не удалось скачать файл.")
        return

    content = _parse_spreadsheet_bytes(data, filename, mime_type)
    if not content:
        await message.answer("❌ Не удалось прочитать файл. Поддерживаются: CSV, XLSX, PDF, DOCX, TXT, JSON.")
        return

    try:
        analysis = await analyze_document(content, question, doc_type)
        await message.answer(f"📊 {analysis}", parse_mode="HTML")
    except Exception:
        logger.exception("File analysis failed for owner %d", owner_id)
        await message.answer("❌ Не удалось проанализировать файл.")


async def _handle_pending_email_edit(
    owner_id: int,
    text: str,
    message: Message,
    session: AsyncSession,
) -> bool:
    """If owner is editing a pending email body, update the draft. Returns True if consumed."""
    pending = _pending_email.get(owner_id)
    if pending is None or pending.get("status") != "edit":
        return False

    await _show_email_preview(
        message, owner_id,
        to=str(pending.get("to") or ""),
        subject=str(pending.get("subject") or ""),
        body=text.strip(),
        has_attachment=bool(pending.get("has_attachment")),
        queued_file_id=pending.get("queued_file_id"),  # type: ignore[arg-type]
        queued_filename=str(pending.get("queued_filename") or "attachment"),
        queued_mime=str(pending.get("queued_mime") or "application/octet-stream"),
    )
    return True


async def _handle_pending_email_address(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
    session: AsyncSession,
) -> bool:
    """If owner is providing an email address for a pending email, handle it. Returns True if consumed."""
    pending = _pending_email.get(owner_id)
    # Skip if no pending, already has "to" (waiting for file), or is in preview state
    if pending is None or "to" in pending or pending.get("status") == "preview":
        return False

    import re
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", text.strip()):
        return False

    email_addr = text.strip()
    recipient_name = str(pending.get("recipient_name") or "")
    contact_id: int | None = pending.get("contact_id")  # type: ignore[assignment]

    # Persist email: update contact record AND save name→email mapping in cfg
    if contact_id is not None:
        await contact_repo.set_email(session, owner_id, contact_id, email_addr)
    if recipient_name:
        await cfg_repo.set_config(session, owner_id, f"email_for:{recipient_name.lower()}", email_addr)

    from services.ai import generate_dispatch_message

    us = await us_repo.get_or_create(session, owner_id)
    body = str(pending.get("literal_body") or "")
    if not body:
        style = await get_style_profile(owner_id, [])
        body = await generate_dispatch_message(
            intent=str(pending.get("body_intent") or ""),
            recipient_name=recipient_name,
            language=us.language,
            style_profile=style,
        )
    subject = str(pending.get("subject") or "Сообщение из NeuroSave")
    queued_file_id: str | None = pending.get("queued_file_id")  # type: ignore[assignment]

    await _show_email_preview(
        message, owner_id, email_addr, subject, body,
        has_attachment=bool(pending.get("has_attachment")),
        queued_file_id=queued_file_id,
        queued_filename=str(pending.get("queued_filename") or "attachment"),
        queued_mime=str(pending.get("queued_mime") or "application/octet-stream"),
    )
    return True


def _extract_attachment_info(message: Message) -> tuple[str | None, str, str]:
    """Return (file_id, filename, mime_type) from a document/photo/video message."""
    if message.document:
        return (
            message.document.file_id,
            message.document.file_name or "document",
            message.document.mime_type or "application/octet-stream",
        )
    if message.photo:
        return message.photo[-1].file_id, "photo.jpg", "image/jpeg"
    if message.video:
        return (
            message.video.file_id,
            message.video.file_name or "video.mp4",
            message.video.mime_type or "video/mp4",
        )
    return None, "attachment", "application/octet-stream"


@router.message(F.document | F.photo | F.video)
async def handle_owner_attachment(message: Message, bot: Bot, session: AsyncSession) -> None:
    if message.from_user is None:
        return
    owner_id = message.from_user.id

    file_id, filename, mime_type = _extract_attachment_info(message)
    if not file_id:
        return

    pending = _pending_email.get(owner_id)

    # Parse caption upfront so we can decide routing before touching the file type
    caption_parsed = None
    has_email_intent = False
    if (pending is None or "to" not in pending) and message.caption:
        caption_parsed = await parse_dispatch_command(message.caption)
        has_email_intent = bool(caption_parsed.is_email and caption_parsed.recipients)

    # Analyzable file (spreadsheet, PDF, DOCX, TXT, JSON):
    # if caption present → analyze immediately; otherwise two-step flow with Cancel button.
    if (
        _is_analyzable_file(filename, mime_type)
        and not has_email_intent
        and not (pending and pending.get("status") in ("attach", "preview"))
    ):
        doc_type = _get_doc_type(filename, mime_type)
        if message.caption:
            await _analyze_uploaded_file(message, bot, owner_id, file_id, filename, mime_type, doc_type=doc_type)
        else:
            _pending_file[owner_id] = {
                "file_id": file_id,
                "filename": filename,
                "mime_type": mime_type,
                "doc_type": doc_type,
                "expires_at": str(time.time() + 600),
            }
            await message.answer(
                "📎 Файл получен. Задайте вопрос по содержимому или напишите «анализ» для общего резюме.\n"
                "<i>Поддерживаются: CSV, XLSX, PDF, DOCX, TXT, JSON</i>",
                parse_mode="HTML",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                    InlineKeyboardButton(text="❌ Отмена", callback_data="file_pending_cancel"),
                ]]),
            )
        return

    # User sent file + caption in one message (no prior pending email)
    if caption_parsed is not None:
        parsed = caption_parsed
        if not has_email_intent:
            return

        from services import gmail as gmail_svc

        gmail_service = await gmail_svc.get_gmail_service(owner_id, session)
        if gmail_service is None:
            await message.answer("❌ Gmail не подключён. Перейди в мини-приложение → Настройки → Интеграции → Gmail.")
            return

        recipient_name = parsed.recipients[0]
        contacts = await contact_repo.find_contacts_by_name(session, owner_id, recipient_name)
        contact = contacts[0] if contacts else None

        if contact is None or not contact.email:
            # Store file_id to download after user provides email
            _pending_email[owner_id] = {
                "recipient_name": recipient_name,
                "subject": parsed.email_subject or "",
                "body_intent": parsed.email_body_intent,
                "literal_body": parsed.email_literal_body,
                "has_attachment": True,
                "contact_id": contact.user_id if contact else None,
                "queued_file_id": file_id,
                "queued_filename": filename,
                "queued_mime": mime_type,
            }
            name_display = contact.name or recipient_name if contact else recipient_name
            await message.answer(f"📧 Для {name_display} нет email-адреса.\nВведи адрес электронной почты:")
            return

        body = parsed.email_literal_body
        if not body:
            settings_obj = await us_repo.get_or_create(session, owner_id)
            lang = settings_obj.language or "ru"
            style = await get_style_profile(owner_id, [])
            body = await generate_dispatch_message(
                intent=parsed.email_body_intent or "",
                recipient_name=contact.name or recipient_name,
                language=lang,
                style_profile=style,
            )
        subject = parsed.email_subject or "Сообщение из NeuroSave"

        try:
            buf = await bot.download(file_id)
            if buf is None:
                await message.answer("❌ Не удалось скачать файл.")
                return
            file_bytes = buf.read()
        except Exception:
            logger.exception("Failed to download attachment for owner %d", owner_id)
            await message.answer("❌ Не удалось скачать файл.")
            return

        try:
            await gmail_svc.send_email(
                gmail_service,
                to=[contact.email],
                subject=subject,
                body=body,
                attachments=[(filename, file_bytes, mime_type)],
            )
            await message.answer(
                f"✅ Письмо с вложением <b>{filename}</b> отправлено на <b>{contact.email}</b>",
                parse_mode="HTML",
            )
        except Exception as exc:
            logger.exception("Gmail send with attachment failed for owner %d: %s", owner_id, exc)
            await message.answer("❌ Не удалось отправить письмо.")
        return

    # Two-step flow: pending email already set up, user sends the file now
    if pending is None or "to" not in pending or pending.get("status") == "preview":
        return

    from services import gmail as gmail_svc

    gmail_service = await gmail_svc.get_gmail_service(owner_id, session)
    if gmail_service is None:
        _pending_email.pop(owner_id, None)
        await message.answer("❌ Gmail не подключён.")
        return

    try:
        buf = await bot.download(file_id)
        if buf is None:
            await message.answer("❌ Не удалось скачать файл.")
            return
        file_bytes = buf.read()
    except Exception:
        logger.exception("Failed to download attachment for owner %d", owner_id)
        await message.answer("❌ Не удалось скачать файл.")
        return

    to = str(pending["to"])
    subject = str(pending.get("subject") or "")
    body = str(pending.get("body") or "")
    _pending_email.pop(owner_id, None)

    try:
        await gmail_svc.send_email(
            gmail_service,
            to=[to],
            subject=subject,
            body=body,
            attachments=[(filename, file_bytes, mime_type)],
        )
        await message.answer(
            f"✅ Письмо с вложением <b>{filename}</b> отправлено на <b>{to}</b>",
            parse_mode="HTML",
        )
    except Exception as exc:
        logger.exception("Gmail send with attachment failed for owner %d: %s", owner_id, exc)
        await message.answer("❌ Не удалось отправить письмо.")


async def _handle_pending_file(
    owner_id: int,
    text: str,
    message: Message,
    bot: Bot,
) -> bool:
    """If owner sent a file earlier and now asks a question, analyze it. Returns True if handled."""
    pf = _pending_file.get(owner_id)
    if pf is None:
        return False
    # If expired, discard without consuming the message so normal handlers run
    if time.time() > float(pf.get("expires_at", 0)):
        _pending_file.pop(owner_id, None)
        return False
    _pending_file.pop(owner_id, None)
    from services.ai import analyze_document
    thinking = await message.answer("🔍 Анализирую файл…")
    try:
        # Re-download the file using stored file_id
        buf = await bot.download(pf["file_id"])
        if buf is None:
            await thinking.edit_text("❌ Не удалось скачать файл.")
            return True
        data = buf.read()
        content = _parse_spreadsheet_bytes(data, pf["filename"], pf["mime_type"])
        if not content:
            await thinking.edit_text("❌ Не удалось прочитать файл. Поддерживаются: CSV, XLSX, PDF, DOCX, TXT, JSON.")
            return True
        if len(content) > 14000:
            content = content[:14000]
        analysis = await analyze_document(content, text, pf["doc_type"])
        await thinking.edit_text(f"📊 {analysis}", parse_mode="HTML")
    except Exception:
        logger.exception("Pending file analysis failed for owner %d", owner_id)
        await thinking.edit_text("❌ Не удалось проанализировать файл.")
    return True


async def _handle_google_doc_url(
    text: str,
    message: Message,
    session: AsyncSession,
    owner_id: int,
) -> bool:
    """Fetch and analyze a Google Docs/Sheets URL pasted by the owner. Returns True if handled."""
    sheets_m = _SHEETS_URL_RE.search(text)
    docs_m = _DOCS_URL_RE.search(text)
    if not sheets_m and not docs_m:
        return False

    if sheets_m:
        doc_id = sheets_m.group(1)
        export_url = f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=csv"
        doc_type = "spreadsheet"
    else:
        doc_id = docs_m.group(1)  # type: ignore[union-attr]
        export_url = f"https://docs.google.com/document/d/{doc_id}/export?format=txt"
        doc_type = "document"

    thinking = await message.answer("🔍 Читаю документ…")

    import httpx
    from db.repositories import oauth as oauth_repo

    content: str | None = None

    # Try authenticated access first (private docs)
    google_token = await oauth_repo.get_token(session, owner_id, "google")
    if google_token:
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(
                    export_url,
                    headers={"Authorization": f"Bearer {google_token.access_token}"},
                    follow_redirects=True,
                    timeout=15,
                )
                if resp.status_code == 200:
                    content = resp.text
        except Exception:
            pass

    # Fallback: public access
    if not content:
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(export_url, follow_redirects=True, timeout=15)
                if resp.status_code == 200:
                    content = resp.text
        except Exception:
            pass

    if not content:
        await thinking.edit_text(
            "❌ Не удалось открыть документ. Убедитесь что он доступен для просмотра, "
            "или подключите Google в мини-приложении → Настройки → Интеграции."
        )
        return True

    if len(content) > 14000:
        content = content[:14000]

    # Use text without the URL as the question
    question = _SHEETS_URL_RE.sub("", _DOCS_URL_RE.sub("", text)).strip()
    if not question:
        question = "Проанализируй содержимое, дай краткую сводку"

    from services.ai import analyze_document
    try:
        analysis = await analyze_document(content, question, doc_type)
        await thinking.edit_text(f"📊 {analysis}", parse_mode="HTML")
    except Exception:
        logger.exception("Google doc analysis failed for owner %d", owner_id)
        await thinking.edit_text("❌ Не удалось проанализировать документ.")
    return True


@router.message(F.text & ~F.text.startswith("/"))
async def handle_owner_dispatch(message: Message, bot: Bot, session: AsyncSession) -> None:
    if message.from_user is None or not message.text:
        return
    owner_id = message.from_user.id
    await us_repo.get_or_create(session, owner_id)
    if await _handle_pending_snooze(owner_id, message.text, message, bot, session):
        return
    if await _handle_pending_ghost_time(owner_id, message.text, message, bot, session):
        return
    if await _handle_pending_ghost_text(owner_id, message.text, message, bot, session):
        return
    if await _handle_pending_gmail_reply(owner_id, message.text, message, bot, session):
        return
    if await _handle_pending_email_edit(owner_id, message.text, message, session):
        return
    if await _handle_pending_email_address(owner_id, message.text, message, bot, session):
        return
    if await _handle_google_doc_url(message.text, message, session, owner_id):
        return
    if await _handle_pending_file(owner_id, message.text, message, bot):
        return
    await _process_owner_text(message.text, message, bot, session, owner_id)


@router.message(F.voice)
async def handle_owner_voice(message: Message, bot: Bot, session: AsyncSession) -> None:
    if message.from_user is None or not message.voice:
        return
    owner_id = message.from_user.id
    thinking = await message.answer("🎙 Распознаю…")
    try:
        buf = await bot.download(message.voice)
        if buf is None:
            await thinking.edit_text("❌ Не удалось скачать голосовое сообщение.")
            return
        text = await transcribe_voice(buf.read())
    except Exception:
        logger.exception("Voice transcription failed")
        await thinking.edit_text("❌ Не удалось распознать голосовое сообщение.")
        return
    await thinking.edit_text(f"🎙 <i>{text}</i>", parse_mode="HTML")
    await us_repo.get_or_create(session, owner_id)
    if await _handle_pending_gmail_reply(owner_id, text, message, bot, session):
        return
    await _process_owner_text(text, message, bot, session, owner_id)


async def cmd_reminders(message: Message, session: AsyncSession) -> None:
    if message.from_user is None:
        return
    owner_id = message.from_user.id
    us = await us_repo.get_or_create(session, owner_id)
    active = await task_repo.get_open_tasks(session, owner_id, is_personal=True)
    if not active:
        await message.answer("Активных напоминаний нет.")
        return
    lines = ["⏰ <b>Активные напоминания:</b>"]
    for i, t in enumerate(active, 1):
        time_str = _format_time_local(_task_reminder_iso(t), us.timezone, us.language)
        lines.append(f"{i}. {t.description}" + (f" — {time_str}" if time_str else ""))
    lines.append("\nЧтобы изменить или удалить — напишите, например: «удали напоминание про стрижку» или «перенеси встречу на 8 утра».")
    await message.answer("\n".join(lines), parse_mode="HTML")
